from __future__ import annotations

import csv
import hashlib
import json
import re
import sqlite3
import sys
import uuid
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = Path(r"E:\www\tmall-dashboard\data\demo\dashboard.db")
CSV_ROOT = Path(r"E:\派米\数据源")
IMAGE_BOOK = Path(r"E:\桌面\0817\标品货品0817.xlsx")
LOG_PATH = ROOT / "data" / "paimi_monthly_import.json"


def clean_text(value):
    if value is None:
        return ""
    return str(value).strip()


def product_id(value):
    text = clean_text(value)
    if not text:
        return ""
    if re.fullmatch(r"\d+\.0+", text):
        return text.split(".", 1)[0]
    return text


def clean_number(value):
    text = clean_text(value)
    if not text or text in {"-", "--", "nan", "None"}:
        return 0.0
    text = text.replace(",", "").replace("¥", "").replace("￥", "")
    try:
        return float(text)
    except ValueError:
        return 0.0


def clean_int(value):
    return int(clean_number(value))


def clean_percentage(value):
    text = clean_text(value)
    if not text or text in {"-", "--", "nan", "None"}:
        return 0.0
    try:
        n = float(text.replace(",", "").replace("%", ""))
    except ValueError:
        return 0.0
    return n / 100.0 if abs(n) > 1 else n


def clean_date(value):
    text = clean_text(value)
    if not text:
        return ""
    if text.startswith('="') and text.endswith('"'):
        text = text[2:-1]
    text = text.strip('"')
    m = re.search(r"(\d{4}-\d{1,2}-\d{1,2})", text)
    return m.group(1) if m else text


def month_from_name(path):
    m = re.search(r"(\d{4}-\d{2})-\d{2}", path.name)
    if not m:
        raise ValueError(f"无法从文件名识别月份: {path.name}")
    return m.group(1)


def valid_image(value):
    text = clean_text(value)
    return text if text.lower().startswith(("http://", "https://")) else ""


def sha256(path):
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def load_bi_images(path):
    wb = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb["BI"]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        idx = {clean_text(v): i for i, v in enumerate(header) if clean_text(v)}
        id_idx = idx["商品id"] if "商品id" in idx else idx.get("商品ID")
        image_idx = idx.get("商品主图")
        title_idx = idx.get("商品")
        status_idx = idx.get("商品状态")
        if id_idx is None or image_idx is None:
            raise RuntimeError(f"BI页缺少商品id/商品主图字段，表头: {list(idx)[:20]}")
        result = {}
        for row in rows:
            if not row:
                continue
            pid = product_id(row[id_idx] if id_idx < len(row) else None)
            image = valid_image(row[image_idx] if image_idx < len(row) else None)
            if pid and image:
                result[pid] = {
                    "image_url": image,
                    "title": clean_text(row[title_idx]) if title_idx is not None and title_idx < len(row) else "",
                    "source_status": clean_text(row[status_idx]) if status_idx is not None and status_idx < len(row) else "",
                }
        return result
    finally:
        wb.close()


def upsert_product(conn, pid, title, category, image_url, list_date, source_status=""):
    conn.execute(
        """
        INSERT INTO products (product_id, title, category, list_date, status, image_url, source_status, updated_at)
        VALUES (?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(product_id) DO UPDATE SET
            title=COALESCE(NULLIF(excluded.title, ''), products.title),
            category=COALESCE(NULLIF(excluded.category, ''), products.category),
            list_date=COALESCE(NULLIF(excluded.list_date, ''), products.list_date),
            image_url=COALESCE(NULLIF(excluded.image_url, ''), products.image_url),
            source_status=COALESCE(NULLIF(excluded.source_status, ''), products.source_status),
            updated_at=CURRENT_TIMESTAMP
        """,
        (pid, title, category, list_date, "active", image_url, source_status),
    )


MONTHLY_SQL = """
INSERT INTO monthly_data (
    product_id, month, payment_amount, refund_amount, net_sales,
    visitors, page_views, uv_value, search_visitors, search_ratio,
    payment_conversion, search_conversion, cart_rate, fav_rate, bounce_rate, avg_stay_duration,
    ad_spend, ad_roi, overall_roi, paid_ratio, refund_paid_ratio,
    keyword_spend, keyword_sales, keyword_roi, keyword_visitors, keyword_ppc,
    crowd_spend, crowd_sales, crowd_roi, crowd_visitors, crowd_ppc,
    site_spend, site_sales, site_roi, site_visitors, site_ppc,
    refund_rate, repurchase_rate, cross_sell_rate,
    buyers, avg_order_value, payment_qty, cart_qty, fav_users, click_rate, score,
    data_source
) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
ON CONFLICT(product_id, month) DO UPDATE SET
    payment_amount=excluded.payment_amount, refund_amount=excluded.refund_amount, net_sales=excluded.net_sales,
    visitors=excluded.visitors, page_views=excluded.page_views, uv_value=excluded.uv_value,
    search_visitors=excluded.search_visitors, search_ratio=excluded.search_ratio,
    payment_conversion=excluded.payment_conversion, search_conversion=excluded.search_conversion,
    cart_rate=excluded.cart_rate, fav_rate=excluded.fav_rate, bounce_rate=excluded.bounce_rate,
    avg_stay_duration=excluded.avg_stay_duration, ad_spend=excluded.ad_spend, ad_roi=excluded.ad_roi,
    overall_roi=excluded.overall_roi, paid_ratio=excluded.paid_ratio, refund_paid_ratio=excluded.refund_paid_ratio,
    keyword_spend=excluded.keyword_spend, keyword_sales=excluded.keyword_sales, keyword_roi=excluded.keyword_roi,
    keyword_visitors=excluded.keyword_visitors, keyword_ppc=excluded.keyword_ppc,
    crowd_spend=excluded.crowd_spend, crowd_sales=excluded.crowd_sales, crowd_roi=excluded.crowd_roi,
    crowd_visitors=excluded.crowd_visitors, crowd_ppc=excluded.crowd_ppc,
    site_spend=excluded.site_spend, site_sales=excluded.site_sales, site_roi=excluded.site_roi,
    site_visitors=excluded.site_visitors, site_ppc=excluded.site_ppc,
    refund_rate=excluded.refund_rate, repurchase_rate=excluded.repurchase_rate, cross_sell_rate=excluded.cross_sell_rate,
    buyers=excluded.buyers, avg_order_value=excluded.avg_order_value, payment_qty=excluded.payment_qty,
    cart_qty=excluded.cart_qty, fav_users=excluded.fav_users, click_rate=excluded.click_rate,
    score=excluded.score, data_source=excluded.data_source
"""


def monthly_values(row, month, source_name):
    def n(k): return clean_number(row.get(k))
    def i(k): return clean_int(row.get(k))
    def p(k): return clean_percentage(row.get(k))
    # 货品全站推广字段在源表中使用“货品全站推广”，而不是旧脚本中的别名。
    return (
        product_id(row.get("商品ID")), month,
        n("支付金额"), n("退款金额"), n("退款后销售额"),
        i("访客数"), i("浏览量"), n("UV价值"), i("搜索人数"), p("搜索占比"),
        p("支付转化率"), p("搜索支付转化率"), p("加购率"), p("访客收藏率"), p("跳失率"), n("平均停留时长"),
        n("总推广花费"), n("推广直接ROI"),
        (n("退款后销售额") / n("总推广花费") if n("总推广花费") else 0.0),
        p("付费占比"), p("退款付费占比"),
        n("关键词推广花费"), n("关键词推广销售额"), n("关键词推广投产"), i("关键词推广访客数"), n("关键词推广PPC"),
        n("人群推广花费"), n("人群推广销售额"), n("人群推广投产"), i("人群推广访客数"), n("人群推广PPC"),
        n("货品全站推广花费"), n("货品全站推广销售额"), n("货品全站推广投产"), i("货品全站推广访客数"), n("货品全站推广PPC"),
        p("退款率"), 0.0, 0.0,
        i("支付人数"), n("客单价"), i("支付件数"), i("加购件数"), i("收藏人数"), p("总点击率"), i("评分"),
        f"paimi-monthly:{month}:{source_name}",
    )


def import_file(conn, path, bi_map):
    month = month_from_name(path)
    batch_id = str(uuid.uuid4())
    file_hash = sha256(path)
    conn.execute(
        "INSERT INTO import_batches (id, source_type, source_filename, source_hash, status, quality_summary) VALUES (?, ?, ?, ?, ?, ?)",
        (batch_id, "product_month", str(path), file_hash, "running", json.dumps({"month": month}, ensure_ascii=False)),
    )
    conn.commit()

    total = valid = invalid = bi_matches = 0
    before = conn.execute("SELECT COUNT(*) FROM monthly_data WHERE month=?", (month,)).fetchone()[0]
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            reader = csv.DictReader(f)
            required = {"商品ID", "商品标题", "图片链接", "商品类目"}
            missing = required - set(reader.fieldnames or [])
            if missing:
                raise RuntimeError(f"{path.name} 缺少字段: {sorted(missing)}")
            for row in reader:
                total += 1
                pid = product_id(row.get("商品ID"))
                if not pid:
                    invalid += 1
                    continue
                valid += 1
                bi = bi_map.get(pid, {})
                if bi.get("image_url"):
                    bi_matches += 1
                image = bi.get("image_url") or valid_image(row.get("图片链接"))
                upsert_product(
                    conn, pid, clean_text(row.get("商品标题")), clean_text(row.get("商品类目")), image,
                    clean_date(row.get("上架时间")), bi.get("source_status", ""),
                )
                conn.execute(MONTHLY_SQL, monthly_values(row, month, path.name))
        after = conn.execute("SELECT COUNT(*) FROM monthly_data WHERE month=?", (month,)).fetchone()[0]
        inserted = max(0, after - before)
        updated = valid - inserted
        quality = {"month": month, "total_rows": total, "valid_rows": valid, "invalid_rows": invalid, "bi_image_matches_seen": bi_matches}
        conn.execute(
            "UPDATE import_batches SET status='completed', total_rows=?, valid_rows=?, invalid_rows=?, inserted_count=?, updated_count=?, quality_summary=?, completed_at=CURRENT_TIMESTAMP WHERE id=?",
            (total, valid, invalid, inserted, max(0, updated), json.dumps(quality, ensure_ascii=False), batch_id),
        )
        conn.commit()
        return {"file": path.name, "month": month, "total": total, "valid": valid, "invalid": invalid, "inserted": inserted, "updated": max(0, updated), "bi_image_matches": bi_matches}
    except Exception:
        conn.rollback()
        conn.execute("UPDATE import_batches SET status='failed', total_rows=?, valid_rows=?, invalid_rows=?, completed_at=CURRENT_TIMESTAMP WHERE id=?", (total, valid, invalid, batch_id))
        conn.commit()
        raise


def main():
    if not DB_PATH.exists(): raise SystemExit(f"数据库不存在: {DB_PATH}")
    files = sorted(CSV_ROOT.glob("*.csv"))
    if not files: raise SystemExit(f"没有找到CSV: {CSV_ROOT}")
    print("读取 BI 图片映射...")
    bi_map = load_bi_images(IMAGE_BOOK)
    print(json.dumps({"bi_products_with_images": len(bi_map)}, ensure_ascii=True))
    conn = sqlite3.connect(DB_PATH, timeout=120)
    conn.execute("PRAGMA foreign_keys=ON")
    conn.execute("PRAGMA busy_timeout=120000")
    results = []
    try:
        for index, path in enumerate(files, 1):
            print(f"[{index}/{len(files)}] 导入 {path.name}", flush=True)
            result = import_file(conn, path, bi_map)
            results.append(result)
            print(json.dumps(result, ensure_ascii=True), flush=True)
    finally:
        conn.close()
    summary = {"generated_at": datetime.now().isoformat(timespec="seconds"), "csv_root": str(CSV_ROOT), "image_book": str(IMAGE_BOOK), "bi_products_with_images": len(bi_map), "files": results, "rows": sum(x["valid"] for x in results)}
    LOG_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print("SUMMARY", json.dumps(summary, ensure_ascii=True))


if __name__ == "__main__":
    main()
