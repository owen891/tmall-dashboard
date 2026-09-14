from __future__ import annotations

import json
import re
import shutil
import sqlite3
import zipfile
from datetime import datetime
from pathlib import Path
import xml.etree.ElementTree as ET

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DB_PATH = ROOT / 'data' / 'demo' / 'dashboard.db'
IMAGE_BOOK = Path(r'E:\桌面\0817\标品货品0817.xlsx')
OUT_DIR = ROOT / 'frontend' / 'ui_demo' / 'assets' / 'product-thumbs'
REPORT_PATH = ROOT / 'data' / 'product_image_repair.json'


def norm_pid(value):
    if value is None:
        return ''
    text = str(value).strip()
    if re.fullmatch(r'\d+\.0+', text):
        text = text.split('.', 1)[0]
    return text


def parse_cell_image_map(zf: zipfile.ZipFile) -> dict[str, str]:
    ns = {
        'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'etc': 'http://www.wps.cn/officeDocument/2017/etCustomData',
        'pr': 'http://schemas.openxmlformats.org/package/2006/relationships',
    }
    root = ET.fromstring(zf.read('xl/cellimages.xml'))
    relroot = ET.fromstring(zf.read('xl/_rels/cellimages.xml.rels'))
    rels = {e.attrib['Id']: e.attrib['Target'] for e in relroot.findall('pr:Relationship', ns)}
    result: dict[str, str] = {}
    for cell_image in root.findall('etc:cellImage', ns):
        name_node = cell_image.find('.//xdr:cNvPr', ns)
        blip_node = cell_image.find('.//a:blip', ns)
        if name_node is None or blip_node is None:
            continue
        image_id = name_node.attrib.get('name', '').strip()
        rid = blip_node.attrib.get('{%s}embed' % ns['r'], '')
        target = rels.get(rid, '').lstrip('/')
        if image_id and target:
            result[image_id] = 'xl/' + target if not target.startswith('xl/') else target
    return result


def parse_data2_map() -> dict[str, str]:
    wb = load_workbook(IMAGE_BOOK, data_only=False, read_only=True, keep_links=False)
    try:
        ws = wb['数据 (2)']
        result: dict[str, str] = {}
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row:
                continue
            pid = norm_pid(row[0] if len(row) > 0 else None)
            formula = str(row[2] if len(row) > 2 else '')
            match = re.search(r'DISPIMG\(\s*["\'](ID_[A-Z0-9]+)["\']', formula, re.I)
            if pid and match:
                result[pid] = match.group(1).upper()
        return result
    finally:
        wb.close()


def parse_bi_urls() -> dict[str, str]:
    wb = load_workbook(IMAGE_BOOK, data_only=True, read_only=True, keep_links=False)
    try:
        ws = wb['BI']
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        indexes = {str(value).strip(): idx for idx, value in enumerate(header) if value is not None}
        id_idx = indexes.get('商品id', indexes.get('商品ID'))
        image_idx = indexes.get('商品主图')
        if id_idx is None or image_idx is None:
            return {}
        result = {}
        for row in rows:
            if len(row) <= max(id_idx, image_idx):
                continue
            pid = norm_pid(row[id_idx])
            image = str(row[image_idx] or '').strip()
            if pid and image.startswith(('http://', 'https://')):
                result[pid] = image
        return result
    finally:
        wb.close()


def main() -> None:
    if not DB_PATH.exists():
        raise SystemExit(f'数据库不存在: {DB_PATH}')
    if not IMAGE_BOOK.exists():
        raise SystemExit(f'Excel 不存在: {IMAGE_BOOK}')

    conn = sqlite3.connect(DB_PATH, timeout=120)
    conn.row_factory = sqlite3.Row
    products = {str(row['product_id']): dict(row) for row in conn.execute('SELECT product_id, image_url FROM products')}
    print(f'products={len(products)}')

    formula_map = parse_data2_map()
    bi_urls = parse_bi_urls()
    print(f'excel_formula_map={len(formula_map)} bi_urls={len(bi_urls)}')

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    extracted = 0
    reused = 0
    missing_media = 0
    local_urls: dict[str, str] = {}
    bad_before = []
    unresolved_bad = []

    with zipfile.ZipFile(IMAGE_BOOK) as zf:
        cell_map = parse_cell_image_map(zf)
        print(f'cell_image_map={len(cell_map)}')
        for index, (pid, current) in enumerate(products.items(), start=1):
            image_id = formula_map.get(pid)
            target = cell_map.get(image_id, '') if image_id else ''
            if target and target in zf.namelist():
                suffix = Path(target).suffix.lower() or '.img'
                out = OUT_DIR / f'{pid}{suffix}'
                if not out.exists() or out.stat().st_size == 0:
                    with zf.open(target) as src, out.open('wb') as dst:
                        shutil.copyfileobj(src, dst, length=1024 * 1024)
                    extracted += 1
                else:
                    reused += 1
                local_urls[pid] = f'/assets/product-thumbs/{out.name}'
            elif image_id:
                missing_media += 1
            url = str(current.get('image_url') or '')
            if re.search(r'/uploaded/\d+\.jpg$', url):
                bad_before.append(pid)

    # If a BI image URL is good but there was no embedded image, use it as a
    # second source. This also repairs malformed numeric URLs from the CSV.
    for pid, url in bi_urls.items():
        if pid in products and pid not in local_urls and url.startswith(('http://', 'https://')):
            local_urls[pid] = url

    for pid in bad_before:
        if pid not in local_urls:
            unresolved_bad.append(pid)

    with conn:
        for pid, url in local_urls.items():
            conn.execute('UPDATE products SET image_url=?, updated_at=CURRENT_TIMESTAMP WHERE product_id=?', (url, pid))
        # Do not keep known-dead numeric URLs: the UI will render its explicit
        # placeholder instead of issuing a guaranteed 404 request.
        for pid in unresolved_bad:
            conn.execute("UPDATE products SET image_url='', updated_at=CURRENT_TIMESTAMP WHERE product_id=?", (pid,))

    report = {
        'generated_at': datetime.now().isoformat(timespec='seconds'),
        'database': str(DB_PATH),
        'workbook': str(IMAGE_BOOK),
        'products': len(products),
        'excel_formula_products': len(formula_map),
        'excel_cell_images': len(cell_map),
        'local_images_extracted': extracted,
        'local_images_reused': reused,
        'local_image_products': len([pid for pid in local_urls if pid in formula_map]),
        'bi_url_fallback_products': len([pid for pid in local_urls if pid not in formula_map]),
        'known_bad_numeric_urls_before': len(bad_before),
        'known_bad_numeric_urls_cleared': len(unresolved_bad),
        'known_bad_numeric_urls_repaired': len(bad_before) - len(unresolved_bad),
        'missing_media_for_formula_products': missing_media,
        'unresolved_product_ids': unresolved_bad,
    }
    REPORT_PATH.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
