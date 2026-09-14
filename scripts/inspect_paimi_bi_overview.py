import argparse
import json
import re
import sqlite3
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

CORE = [
    '商品id','商品','商品主图','商品标签','渠道','店铺','商品状态','上架时间',
    '净销售额(支付)','支付金额(支付)','支付金额(发货)','净销售额(发货)',
    '普通单退款金额','退款金额','普通单退款率(按金额)','支付人数','支付转化率',
    '销售件数(支付)','销售单数(支付)','加购人数','加购率','收藏人数','收藏率',
    '访客数','浏览量','访客价值','搜索访客数','搜索访客占比','搜索点击率','搜索转化率',
    '推广花费(账单)','推广花费(支付预估)','推广花费(发货预估)',
    '利润(支付预估)','利润(账单)','毛利率','商品分类','商品类目','负责人',
]

def clean(v):
    if v is None: return None
    if isinstance(v, str):
        s=v.strip()
        return s if s not in {'','-','--','—','－'} else None
    return v

def number(v):
    v=clean(v)
    if v is None: return None
    if isinstance(v,(int,float)): return float(v)
    s=str(v).replace(',','').replace('￥','').replace('¥','').strip()
    if s.endswith('%'):
        try: return float(s[:-1])/100
        except ValueError: return None
    try: return float(s)
    except ValueError: return None

def row_is_data(pid):
    s=str(pid or '').strip()
    if not s: return False
    return s.lower() not in {'合计','总计','汇总','nan','none'}

def scan(path):
    wb=load_workbook(path, read_only=True, data_only=True)
    ws=wb[wb.sheetnames[0]]
    it=ws.iter_rows(values_only=True)
    h=list(next(it)); _=next(it, None)
    idx={str(v).strip():i for i,v in enumerate(h) if v is not None and str(v).strip()}
    missing=[x for x in CORE if x not in idx]
    counts=Counter(); ids=Counter(); sums=defaultdict(float); nonempty=Counter(); samples=[]
    raw_rows=0; data_rows=0; invalid=[]
    for row in it:
        raw_rows += 1
        pid=clean(row[idx['商品id']]) if '商品id' in idx and idx['商品id'] < len(row) else None
        if not row_is_data(pid):
            invalid.append({'row':raw_rows+2,'product_id':pid,'title':clean(row[idx['商品']]) if '商品' in idx and idx['商品'] < len(row) else None})
            continue
        data_rows += 1
        pid=str(pid)
        ids[pid]+=1
        if len(samples)<3: samples.append({k: clean(row[i]) if i < len(row) else None for k,i in idx.items() if k in CORE})
        for k in CORE:
            if k in idx and idx[k] < len(row):
                v=clean(row[idx[k]])
                if v is not None: nonempty[k]+=1
                n=number(v)
                if n is not None and k in {'净销售额(支付)','支付金额(支付)','支付金额(发货)','净销售额(发货)','普通单退款金额','退款金额','支付人数','销售件数(支付)','销售单数(支付)','加购人数','收藏人数','访客数','浏览量','访客价值','搜索访客数','推广花费(账单)','推广花费(支付预估)','推广花费(发货预估)','利润(支付预估)','利润(账单)'}:
                    sums[k]+=n
    wb.close()
    duplicates={k:v for k,v in ids.items() if v>1}
    return {
        'file': path.name, 'month': re.search(r'(20\d\d-\d\d)',path.name).group(1) if re.search(r'(20\d\d-\d\d)',path.name) else '',
        'raw_rows':raw_rows,'data_rows':data_rows,'unique_product_ids':len(ids),'duplicate_ids_count':len(duplicates),
        'duplicate_examples':dict(list(duplicates.items())[:20]),'excluded_rows':invalid[:20],
        'missing_core':missing,'nonempty_counts':dict(nonempty),'sums':dict(sums),'samples':samples,
    }

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--source', default=r'E:\派米\BI数据源25.4-26.7'); ap.add_argument('--db', default=r'E:\www\tmall-dashboard\data\demo\dashboard.db'); ap.add_argument('--out', default='tmp_paimi_bi_inspection.json'); args=ap.parse_args()
    results=[scan(f) for f in sorted(Path(args.source).glob('*.xlsx'))]
    compare=[]
    con=sqlite3.connect(args.db); con.row_factory=sqlite3.Row
    for r in results:
        db=con.execute('''SELECT COUNT(*) rows, COUNT(DISTINCT product_id) unique_product_ids,
          SUM(payment_amount) payment_amount,SUM(refund_amount) refund_amount,SUM(net_sales) net_sales,
          SUM(visitors) visitors,SUM(buyers) buyers,SUM(ad_spend) ad_spend FROM monthly_data WHERE month=?''',(r['month'],)).fetchone()
        r['db']={k:db[k] for k in db.keys()}
        compare.append({'month':r['month'],'source_rows':r['data_rows'],'source_ids':r['unique_product_ids'],'db_rows':db['rows'],'db_ids':db['unique_product_ids'], 'source_payment':r['sums'].get('支付金额(支付)'), 'db_payment':db['payment_amount'], 'source_net_sales':r['sums'].get('净销售额(支付)'), 'db_net_sales':db['net_sales'], 'source_visitors':r['sums'].get('访客数'), 'db_visitors':db['visitors'], 'source_buyers':r['sums'].get('支付人数'), 'db_buyers':db['buyers'], 'source_ad_spend':r['sums'].get('推广花费(账单)'), 'db_ad_spend':db['ad_spend']})
    con.close()
    payload={'source':str(Path(args.source)),'results':results,'compare':compare}
    out=Path(args.out); out.write_text(json.dumps(payload,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
    print(out.resolve())
    for x in compare: print(json.dumps(x,ensure_ascii=False))
    for r in results:
        if r['duplicate_ids_count'] or r['excluded_rows']: print('ISSUE',r['month'],r['duplicate_ids_count'],r['excluded_rows'][:3])
if __name__=='__main__': main()
