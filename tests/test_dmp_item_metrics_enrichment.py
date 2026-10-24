"""达摩盘单品指标回填：只补 NULL、窗口对齐才生效、不得撑错导出列数。

背景：`daily_data` 里 search_click_rate / category_width / repurchase_* /
cross_sell_* / presale_* 这些列从来没有出现过非零值（采集链路没落库），
唯一真实源是达摩盘采集器写入的 `dmp_shop_item_indicators`。这些指标按 7/30 天
滚动窗口快照存在，不能直接写进日粒度事实表，所以采用查询期回填。
"""

import os
import sys
import unittest
from io import BytesIO

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

DMP_DDL = '''
CREATE TABLE IF NOT EXISTS dmp_shop_item_indicators (
    id INTEGER PRIMARY KEY AUTOINCREMENT,
    run_id INTEGER,
    end_date TEXT,
    period INTEGER,
    item_id TEXT,
    ind_key TEXT,
    ind_name TEXT,
    parent_key TEXT,
    value REAL,
    peer_value REAL,
    generalized_peer TEXT,
    growth_rate REAL,
    created_at TEXT
)
'''


def _insert_run(conn, run_id, period, end_date, items):
    """items: {item_id: {ind_key: value}}"""
    for item_id, metrics in items.items():
        for ind_key, value in metrics.items():
            conn.execute(
                '''INSERT INTO dmp_shop_item_indicators
                   (run_id, end_date, period, item_id, ind_key, value)
                   VALUES (?,?,?,?,?,?)''',
                (run_id, end_date, period, str(item_id), ind_key, value),
            )


def _insert_peer(conn, run_id, period, end_date, items):
    """items: {item_id: {ind_key: peer_value}} —— 只写同行同层 peer_value。"""
    for item_id, metrics in items.items():
        for ind_key, peer in metrics.items():
            conn.execute(
                '''INSERT INTO dmp_shop_item_indicators
                   (run_id, end_date, period, item_id, ind_key, peer_value)
                   VALUES (?,?,?,?,?,?)''',
                (run_id, end_date, period, str(item_id), ind_key, peer),
            )


class DmpSnapshotResolverTests(unittest.TestCase):
    def setUp(self):
        import tempfile
        from app import create_app
        from db import get_db

        self.tmp = tempfile.TemporaryDirectory(prefix='tmall-dmp-enrich-')
        self.db = os.path.join(self.tmp.name, 'dashboard.db')
        self.app = create_app({'TESTING': True, 'DATABASE_PATH': self.db})
        self.client = self.app.test_client()
        with get_db(self.db) as c:
            c.execute(DMP_DDL)
            _insert_run(c, 5, 7, '2026-10-18', {
                'p-a': {'searchCtr': 0.05, 'relatedAlipayCateCnt': 3, 'rebuyRate': 0.2,
                        'rebuyUv': 40, 'relatedAlipayCnt': 120, 'relatedAlipayRate': 0.15},
            })
            _insert_run(c, 5, 30, '2026-10-18', {
                'p-a': {'searchCtr': 0.07, 'rebuyRate': 0.1},
            })
            c.commit()

    def tearDown(self):
        self.tmp.cleanup()

    def _conn(self):
        from db import get_db
        return get_db(self.db)

    def test_resolves_exact_seven_day_window(self):
        from services.dmp_item_metrics import resolve_dmp_snapshot
        with self._conn() as c:
            info = resolve_dmp_snapshot(c, '2026-10-12', '2026-10-18')
        self.assertIsNotNone(info)
        self.assertEqual(info['run_id'], 5)
        self.assertEqual(info['period'], 7)
        self.assertEqual(info['end_date'], '2026-10-18')
        self.assertEqual(info['aligned_days'], 0)
        self.assertTrue(info['exact'])

    def test_resolves_thirty_day_window(self):
        from services.dmp_item_metrics import resolve_dmp_snapshot
        with self._conn() as c:
            info = resolve_dmp_snapshot(c, '2026-09-19', '2026-10-18')
        self.assertIsNotNone(info)
        self.assertEqual(info['period'], 30)

    def test_rejects_window_without_matching_period(self):
        """18 天区间没有对应快照：宁可返回 None，也不拿 7/30 天数据冒充。"""
        from services.dmp_item_metrics import resolve_dmp_snapshot
        with self._conn() as c:
            self.assertIsNone(resolve_dmp_snapshot(c, '2026-10-01', '2026-10-18'))
            self.assertIsNone(resolve_dmp_snapshot(c, '2026-10-18', '2026-10-18'))

    def test_discloses_alignment_gap_for_recent_snapshot(self):
        from services.dmp_item_metrics import resolve_dmp_snapshot
        with self._conn() as c:
            info = resolve_dmp_snapshot(c, '2026-10-13', '2026-10-19')
        self.assertIsNotNone(info)
        self.assertEqual(info['aligned_days'], 1)
        self.assertFalse(info['exact'])


class DmpEnrichmentTests(unittest.TestCase):
    def setUp(self):
        import tempfile
        from app import create_app
        from db import get_db

        self.tmp = tempfile.TemporaryDirectory(prefix='tmall-dmp-enrich-rows-')
        self.db = os.path.join(self.tmp.name, 'dashboard.db')
        self.app = create_app({'TESTING': True, 'DATABASE_PATH': self.db})
        self.client = self.app.test_client()
        with get_db(self.db) as c:
            c.execute(DMP_DDL)
            c.execute('INSERT INTO products (product_id,title,status) VALUES (?,?,?)',
                      ('p-a', 'A', 'active'))
            c.execute('INSERT INTO products (product_id,title,status) VALUES (?,?,?)',
                      ('p-b', 'B', 'active'))
            # p-a 有一列已有主源值（cart_rate），其余目标列为 NULL。
            c.execute('''INSERT INTO daily_data
                (product_id, date, payment_amount, ipv, cart_rate)
                VALUES (?,?,?,?,?)''', ('p-a', '2026-10-18', 500.0, 900, 0.12))
            # p-b 的 search_click_rate 已有主源真值（0.99），不能被达摩盘覆盖。
            c.execute('''INSERT INTO daily_data
                (product_id, date, payment_amount, ipv, cart_rate, search_click_rate)
                VALUES (?,?,?,?,?,?)''', ('p-b', '2026-10-18', 300.0, 400, 0.05, 0.99))
            _insert_run(c, 5, 7, '2026-10-18', {
                'p-a': {'searchCtr': 0.05, 'relatedAlipayCateCnt': 3, 'rebuyRate': 0.2,
                        'rebuyUv': 40, 'relatedAlipayCnt': 120, 'relatedAlipayRate': 0.15},
                'p-b': {'searchCtr': 0.03, 'relatedAlipayCateCnt': 2, 'rebuyRate': 0.1,
                        'rebuyUv': 8, 'relatedAlipayCnt': 30, 'relatedAlipayRate': 0.09},
            })
            # 行业搜索点击率取 searchCtr 的 peer_value（同行同层），不是本店 value。
            _insert_peer(c, 5, 7, '2026-10-18', {
                'p-a': {'searchCtr': 0.11},
                'p-b': {'searchCtr': 0.09},
            })
            c.commit()

    def tearDown(self):
        self.tmp.cleanup()

    def test_fills_null_metrics_from_snapshot(self):
        response = self.client.get(
            '/api/products?dim=daily&start=2026-10-12&end=2026-10-18&status=all')
        self.assertEqual(response.status_code, 200)
        rows = {row['product_id']: row for row in response.get_json()['data']['rows']}
        row = rows['p-a']
        self.assertAlmostEqual(row['search_click_rate'], 0.05)
        self.assertEqual(row['category_width'], 3)
        self.assertAlmostEqual(row['repurchase_rate'], 0.2)
        self.assertEqual(row['repurchase_users'], 40)
        self.assertEqual(row['cross_sell_qty'], 120)
        self.assertAlmostEqual(row['cross_sell_rate'], 0.15)
        self.assertAlmostEqual(row['industry_ctr'], 0.11)

    def test_industry_ctr_uses_peer_value_not_own_value(self):
        """行业搜索点击率必须取 peer_value，不能误用本店 searchCtr 的 value。"""
        response = self.client.get(
            '/api/products?dim=daily&start=2026-10-12&end=2026-10-18&status=all')
        rows = {row['product_id']: row for row in response.get_json()['data']['rows']}
        self.assertAlmostEqual(rows['p-b']['industry_ctr'], 0.09)
        self.assertAlmostEqual(rows['p-b']['search_click_rate'], 0.99)

    def test_primary_value_is_not_overwritten(self):
        """主源已有真值时，达摩盘不得覆盖（0 是合法主源值同理）。"""
        response = self.client.get(
            '/api/products?dim=daily&start=2026-10-12&end=2026-10-18&status=all')
        rows = {row['product_id']: row for row in response.get_json()['data']['rows']}
        self.assertAlmostEqual(rows['p-b']['search_click_rate'], 0.99)

    def test_discloses_snapshot_in_source_batches(self):
        response = self.client.get(
            '/api/products?dim=daily&start=2026-10-12&end=2026-10-18&status=all')
        payload = response.get_json()
        batches = payload.get('source_batches')
        self.assertTrue(batches, '必须披露达摩盘快照窗口')
        batch = batches[0]
        self.assertEqual(batch['source'], 'dmp_shop_item_indicators')
        self.assertEqual(batch['snapshot_end_date'], '2026-10-18')
        self.assertEqual(batch['period_days'], 7)
        self.assertTrue(any('达摩盘' in note for note in payload.get('limitations') or []))

    def test_no_enrichment_when_window_has_no_snapshot(self):
        response = self.client.get(
            '/api/products?dim=daily&start=2026-10-01&end=2026-10-18&status=all')
        self.assertEqual(response.status_code, 200)
        row = response.get_json()['data']['rows'][0]
        self.assertIsNone(row['search_click_rate'])
        self.assertFalse(response.get_json().get('source_batches'))

    def test_enrichment_does_not_add_columns_to_export(self):
        """导出只挑 2 列时，回填不得凭空新增键（否则表头与数据错位）。"""
        from openpyxl import load_workbook
        columns = ['product_id', 'search_click_rate']
        response = self.client.post('/api/export', json={
            'dim': 'daily', 'start': '2026-10-12', 'end': '2026-10-18',
            'status': 'all', 'columns': columns,
        })
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.data)).active
        header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        data_row = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))
        self.assertEqual(len(header), len(columns))
        self.assertEqual(len(data_row), len(columns))
        self.assertAlmostEqual(data_row[1], 0.05)


if __name__ == '__main__':
    unittest.main(verbosity=2)
