import os
import tempfile
import unittest
from io import BytesIO

from openpyxl import load_workbook


class ProductOverviewApiTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='tmall-product-overview-')
        from app import create_app
        from db import get_db

        self.db = os.path.join(self.temp.name, 'dashboard.db')
        self.app = create_app({'TESTING': True, 'DATABASE_PATH': self.db})
        self.client = self.app.test_client()
        with get_db(self.db) as conn:
            conn.executemany(
                'INSERT INTO products (product_id, title, status, shop_label) VALUES (?, ?, ?, ?)',
                [('risk-product', '高风险商品', 'active', '标品袜子'), ('healthy-product', '健康商品', 'active', '标品袜子')],
            )
            conn.executemany(
                '''INSERT INTO monthly_data
                   (product_id, month, payment_amount, refund_amount, visitors,
                    payment_conversion, cart_rate, ad_spend, avg_order_value)
                   VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)''',
                [
                    ('risk-product', '2026-04', 100, 30, 200, 0.01, 0.12, 40, 100),
                    ('healthy-product', '2026-04', 100, 5, 200, 0.05, 0.03, 10, 300),
                ],
            )
            conn.execute(
                'INSERT INTO products (product_id, title, status, shop_label) VALUES (?, ?, ?, ?)',
                ('catalog-only', 'catalog-only', 'active', '鏍囧搧琚滃瓙'),
            )
            conn.commit()

    def tearDown(self):
        self.temp.cleanup()

    def test_rule_results_are_grouped_and_auditable(self):
        response = self.client.get('/api/product-overview/rule-results?period=2026-04&category_mode=sock')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertTrue(payload['ok'])
        self.assertEqual(payload['data']['eligible_count'], 2)

        rules = {rule['key']: rule for rule in payload['data']['rules']}
        self.assertEqual(rules['high_paid_high_refund']['matched_count'], 1)
        self.assertEqual(rules['high_paid_high_refund']['matched_product_ids'], ['risk-product'])
        self.assertEqual(rules['high_aov']['matched_product_ids'], ['healthy-product'])
        self.assertEqual(rules['high_paid_high_refund']['period'], '2026-04')
        self.assertEqual(payload['evidence'][0]['source'], 'monthly_data')
        self.assertEqual(payload['data']['scope']['shop_label'], '标品袜子')
        self.assertEqual(rules['high_paid_high_refund']['version'], '2026.09')
        self.assertEqual(rules['high_paid_high_refund']['rule_type'], 'diagnosis')
        self.assertEqual(rules['high_paid_high_refund']['conditions'][0]['field'], 'expense_ratio')
        self.assertEqual(rules['high_paid_high_refund']['evidence_summary']['grain'], 'product-month')

    def test_rule_registry_exposes_structured_audit_definition(self):
        response = self.client.get('/api/product-overview/rules')
        self.assertEqual(response.status_code, 200)
        payload = response.get_json(); response.close()
        self.assertEqual(payload['data']['rule_version'], '2026.09')
        rule = next(item for item in payload['data']['rules'] if item['key'] == 'low_conversion')
        self.assertEqual(rule['rule_type'], 'alert')
        self.assertEqual(rule['conditions'][0]['operator'], '>=')

    def test_product_list_uses_same_rule_condition_and_count(self):
        response = self.client.get(
            '/api/products?dim=monthly&period=2026-04&status=active&category_mode=sock&rule_key=high_paid_high_refund&limit=1'
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload['data']['total'], 1)
        self.assertEqual(len(payload['data']['rows']), 1)
        self.assertEqual(payload['data']['rows'][0]['product_id'], 'risk-product')

    def test_scoped_product_list_excludes_catalog_only_rows(self):
        response = self.client.get(
            '/api/products?dim=monthly&period=2026-04&status=active&category_mode=sock&rule_key=all_products&limit=20'
        )
        self.assertEqual(response.status_code, 200)
        payload = response.get_json()
        self.assertEqual(payload['data']['total'], 2)
        self.assertEqual(
            {row['product_id'] for row in payload['data']['rows']},
            {'risk-product', 'healthy-product'},
        )

    def test_export_uses_the_same_rule_scope(self):
        response = self.client.post(
            '/api/export',
            json={
                'type': 'products',
                'dim': 'monthly',
                'period': '2026-04',
                'status': 'active',
                'category_mode': 'sock',
                'rule_key': 'high_paid_high_refund',
                'columns': ['product_id'],
            },
        )
        self.assertEqual(response.status_code, 200)
        sheet = load_workbook(BytesIO(response.data), read_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(len(rows) - 1, 1)
        self.assertEqual(rows[1][0], 'risk-product')
        context = load_workbook(BytesIO(response.data), read_only=True)['查询上下文']
        context_rows = {row[0]: row[1] for row in context.iter_rows(values_only=True) if row[0]}
        self.assertEqual(context_rows['规则键'], 'high_paid_high_refund')
        self.assertEqual(context_rows['规则版本'], '2026.09')
        self.assertEqual(context_rows['商品结果数'], 1)

    def test_unknown_product_rule_is_rejected(self):
        response = self.client.get('/api/products?dim=monthly&period=2026-04&rule_key=missing-rule')
        self.assertEqual(response.status_code, 422)
        self.assertEqual(response.get_json()['code'], 'VALIDATION_ERROR')


if __name__ == '__main__':
    unittest.main(verbosity=2)
